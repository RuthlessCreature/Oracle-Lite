import json
from pathlib import Path

from docx import Document
from PIL import Image

from oracle_lite.config import AppConfig
from oracle_lite.dataset import build_domain_dataset
from oracle_lite.ingest import ingest_corpus
from oracle_lite.scanner import scan_corpus
from oracle_lite.snapshot import create_snapshot


def test_multi_image_segment_is_split_to_one_image_per_training_record(tmp_path: Path):
    cfg = AppConfig(
        minimax_api_key="sk-cp-test",
        corpus_dir=tmp_path / "corpus",
        output_dir=tmp_path / "output",
    )
    cfg.ensure_dirs()

    image_a = tmp_path / "a.png"
    image_b = tmp_path / "b.png"
    Image.new("RGB", (32, 32), "white").save(image_a)
    Image.new("RGB", (32, 32), "black").save(image_b)

    source = cfg.corpus_dir / "manual.docx"
    doc = Document()
    doc.add_paragraph("Bearing inspection reference")
    doc.add_picture(str(image_a))
    doc.add_picture(str(image_b))
    doc.save(source)

    scan_corpus(cfg, verify_all=True)
    stats = ingest_corpus(cfg)
    assert stats.failed == 0

    snap = create_snapshot(cfg, name="dataset-memory-test", mode="full")
    dataset = build_domain_dataset(cfg, snapshot_id=snap.snapshot_id)

    records = []
    for shard in dataset.shard_paths:
        with shard.open("r", encoding="utf-8") as f:
            records.extend(json.loads(line) for line in f if line.strip())

    visual = [
        row for row in records
        if row["mode"] == "multimodal" and row["source_path"].endswith("manual.docx")
    ]
    assert len(visual) == 2
    assert all(len(row["images"]) == 1 for row in visual)


def _register_snapshot_source(cfg, source):
    import hashlib
    from oracle_lite.canonical import CanonicalDocument
    from oracle_lite.db import Registry

    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    registry = Registry(cfg.registry_path)
    registry.upsert_source(
        root=cfg.corpus_dir, path=source, size=source.stat().st_size,
        mtime_ns=source.stat().st_mtime_ns, content_hash=digest,
    )
    canonical = cfg.canonical_dir / f"{digest}.json"
    CanonicalDocument(
        content_hash=digest, source_path=str(source), source_type=source.suffix,
        parser_name="fixture", parser_version=cfg.parser_version,
        text="Original source content",
    ).write_json(canonical)
    registry.save_artifact(
        content_hash=digest, parser_version=cfg.parser_version,
        canonical_path=str(canonical), status="ready",
    )


def test_qa_dataset_uses_frozen_snapshot_after_live_file_deleted(tmp_path):
    cfg = AppConfig(minimax_api_key="test", corpus_dir=tmp_path / "corpus",
                    output_dir=tmp_path / "output")
    cfg.ensure_dirs()
    qa = cfg.corpus_dir / "qa_pairs.jsonl"
    qa.write_text(json.dumps({"question": "What is the pressure limit?",
                              "answer": "12 kPa", "context": "Pressure limit: 12 kPa"}) + "\n")
    _register_snapshot_source(cfg, qa)
    snapshot = create_snapshot(cfg, name="frozen-qa")
    qa.unlink()
    dataset = build_domain_dataset(cfg, snapshot_id=snapshot.snapshot_id)
    rows = [json.loads(line) for path in dataset.shard_paths
            for line in path.read_text().splitlines()]
    assert dataset.qa_records == 1
    assert rows[0]["answer"] == "12 kPa"
    assert rows[0]["content_hash"]
    assert rows[0]["role"] == "current"


def test_qa_added_after_snapshot_is_not_used(tmp_path):
    cfg = AppConfig(minimax_api_key="test", corpus_dir=tmp_path / "corpus",
                    output_dir=tmp_path / "output")
    cfg.ensure_dirs()
    source = cfg.corpus_dir / "spec.txt"
    source.write_text("Pressure limit: 12 kPa")
    _register_snapshot_source(cfg, source)
    snapshot = create_snapshot(cfg, name="without-qa")
    (cfg.corpus_dir / "qa_pairs.jsonl").write_text(
        json.dumps({"question": "New?", "answer": "90 kPa", "context": "90 kPa"}) + "\n"
    )
    dataset = build_domain_dataset(cfg, snapshot_id=snapshot.snapshot_id)
    assert dataset.qa_records == 0


def test_snapshot_rejects_qa_modified_after_scan(tmp_path):
    import pytest
    cfg = AppConfig(minimax_api_key="test", corpus_dir=tmp_path / "corpus",
                    output_dir=tmp_path / "output")
    cfg.ensure_dirs()
    qa = cfg.corpus_dir / "qa_pairs.jsonl"
    qa.write_text(json.dumps({"question": "Limit?", "answer": "12", "context": "12"}) + "\n")
    _register_snapshot_source(cfg, qa)
    qa.write_text(json.dumps({"question": "Limit?", "answer": "90", "context": "90"}) + "\n")
    with pytest.raises(ValueError, match="changed after scan"):
        create_snapshot(cfg, name="stale-qa")
