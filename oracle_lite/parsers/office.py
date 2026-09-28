from __future__ import annotations

import zipfile
from pathlib import Path

import fitz
from docx import Document
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .base import ParsedDocument, ParsedSegment


def _safe_suffix(name: str, fallback: str = ".png") -> str:
    suffix = Path(name).suffix.lower()
    return suffix if suffix else fallback


def _extract_zip_media(path: Path, asset_dir: Path, prefix: str) -> list[str]:
    asset_dir.mkdir(parents=True, exist_ok=True)
    assets: list[str] = []
    with zipfile.ZipFile(path) as zf:
        names = sorted(name for name in zf.namelist() if name.startswith(prefix) and not name.endswith("/"))
        for idx, name in enumerate(names, start=1):
            target = asset_dir / f"media-{idx:04d}{_safe_suffix(name)}"
            target.write_bytes(zf.read(name))
            assets.append(str(target.resolve()))
    return assets


def parse_docx(path: Path, asset_dir: Path) -> ParsedDocument:
    doc = Document(path)
    blocks: list[str] = []
    headings = 0
    tables = 0

    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "").lower() if p.style else ""
        if style.startswith("heading"):
            level = 1
            parts = style.split()
            if parts and parts[-1].isdigit():
                level = max(1, min(6, int(parts[-1])))
            blocks.append(f"{'#' * level} {text}")
            headings += 1
        else:
            blocks.append(text)

    for table in doc.tables:
        rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells] for row in table.rows]
        if not rows:
            continue
        width = max(len(r) for r in rows)
        rows = [r + [""] * (width - len(r)) for r in rows]
        blocks.append("| " + " | ".join(rows[0]) + " |")
        blocks.append("| " + " | ".join(["---"] * width) + " |")
        for row in rows[1:]:
            blocks.append("| " + " | ".join(row) + " |")
        tables += 1

    text = "\n\n".join(blocks)
    images = _extract_zip_media(path, asset_dir, "word/media/")
    return ParsedDocument(
        title=path.stem,
        text=text,
        segments=[ParsedSegment(text=text, images=images, metadata={"kind": "docx"})],
        metadata={
            "extension": ".docx",
            "headings": headings,
            "tables": tables,
            "images": len(images),
        },
    )


def parse_pdf(path: Path, asset_dir: Path) -> ParsedDocument:
    asset_dir.mkdir(parents=True, exist_ok=True)
    doc = fitz.open(path)
    segments: list[ParsedSegment] = []
    pages: list[str] = []
    empty_pages = 0

    for idx, page in enumerate(doc, start=1):
        text = page.get_text("text").strip()
        if not text:
            empty_pages += 1

        # Preserve the actual page as a visual asset. 1.5x is a compromise for
        # RTX 4080 training: enough detail for diagrams/tables without exploding VRAM.
        pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
        image_path = asset_dir / f"page-{idx:04d}.png"
        pix.save(str(image_path))

        page_text = f"## Page {idx}\n\n{text}" if text else f"## Page {idx}"
        pages.append(page_text)
        segments.append(
            ParsedSegment(
                text=text,
                images=[str(image_path.resolve())],
                metadata={"kind": "pdf_page", "page": idx},
            )
        )

    doc.close()
    return ParsedDocument(
        title=path.stem,
        text="\n\n".join(pages),
        segments=segments,
        metadata={
            "extension": ".pdf",
            "pages": len(segments),
            "empty_pages": empty_pages,
            "images": len(segments),
        },
    )


def parse_pptx(path: Path, asset_dir: Path) -> ParsedDocument:
    prs = Presentation(path)
    asset_dir.mkdir(parents=True, exist_ok=True)
    segments: list[ParsedSegment] = []
    slide_blocks: list[str] = []
    table_count = 0
    image_count = 0

    for slide_idx, slide in enumerate(prs.slides, start=1):
        items: list[tuple[int, int, str]] = []
        slide_images: list[str] = []

        for shape_idx, shape in enumerate(slide.shapes, start=1):
            top = int(getattr(shape, "top", 0))
            left = int(getattr(shape, "left", 0))

            if getattr(shape, "has_text_frame", False):
                text = "\n".join(
                    p.text.strip()
                    for p in shape.text_frame.paragraphs
                    if p.text.strip()
                )
                if text:
                    items.append((top, left, text))

            if getattr(shape, "has_table", False):
                rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells] for row in shape.table.rows]
                if rows:
                    width = max(len(r) for r in rows)
                    rows = [r + [""] * (width - len(r)) for r in rows]
                    md = [
                        "| " + " | ".join(rows[0]) + " |",
                        "| " + " | ".join(["---"] * width) + " |",
                    ]
                    md.extend("| " + " | ".join(r) + " |" for r in rows[1:])
                    items.append((top, left, "\n".join(md)))
                    table_count += 1

            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                ext = _safe_suffix(shape.image.filename, ".png")
                target = asset_dir / f"slide-{slide_idx:04d}-image-{shape_idx:03d}{ext}"
                target.write_bytes(shape.image.blob)
                slide_images.append(str(target.resolve()))
                image_count += 1

        items.sort(key=lambda x: (x[0], x[1]))
        body = "\n\n".join(text for _, _, text in items)
        slide_blocks.append(f"## Slide {slide_idx}\n\n{body}".rstrip())
        segments.append(
            ParsedSegment(
                text=body,
                images=slide_images,
                metadata={"kind": "pptx_slide", "slide": slide_idx},
            )
        )

    return ParsedDocument(
        title=path.stem,
        text="\n\n".join(slide_blocks),
        segments=segments,
        metadata={
            "extension": ".pptx",
            "slides": len(prs.slides),
            "tables": table_count,
            "images": image_count,
        },
    )
